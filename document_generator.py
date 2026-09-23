"""Vervangt alleen Word-placeholders en exporteert optioneel met Word naar PDF."""

from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement


PLACEHOLDER = re.compile(r"\{\{([A-Z][A-Z0-9_]*)\}\}")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


class TemplateError(Exception):
    """Het sjabloon bevat een niet ingevulde of ongeldige placeholder."""


def _set_text(node, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set(XML_SPACE, "preserve")
    else:
        node.attrib.pop(XML_SPACE, None)


def _insert_replacement(node, value: str) -> None:
    """Voeg eventuele nieuwe regels in als Word-regelafbrekingen in dezelfde run."""
    parent = node.getparent()
    position = parent.index(node) + 1
    lines = value.split("\n")
    for line in lines[1:]:
        br = OxmlElement("w:br")
        parent.insert(position, br)
        position += 1
        text = OxmlElement("w:t")
        _set_text(text, line)
        parent.insert(position, text)
        position += 1


def _replace_in_paragraph(paragraph, values: dict[str, str], unknown: set[str]) -> None:
    # Werk op tekstknooppunten: placeholders kunnen over meerdere opmaakruns
    # verdeeld zijn. Overige XML (logo's, velden, lijnen) blijft onaangeroerd.
    nodes = paragraph._p.xpath(".//w:t")
    if not nodes:
        return
    chunks = [node.text or "" for node in nodes]
    joined = "".join(chunks)
    offsets = []
    for index, chunk in enumerate(chunks):
        offsets.extend((index, pos) for pos in range(len(chunk)))

    for match in reversed(list(PLACEHOLDER.finditer(joined))):
        key = match.group(1)
        if key not in values:
            unknown.add(key)
            continue
        start_node, start_pos = offsets[match.start()]
        end_node, end_pos = offsets[match.end() - 1]
        first = nodes[start_node]
        last = nodes[end_node]
        prefix = (first.text or "")[:start_pos]
        suffix = (last.text or "")[end_pos + 1:]
        replacement = values[key].replace("\r\n", "\n").replace("\r", "\n")
        lines = replacement.split("\n")

        if first is last:
            _set_text(first, prefix + lines[0] + (suffix if len(lines) == 1 else ""))
            if len(lines) > 1:
                _insert_replacement(first, replacement)
                # De suffix moet na de laatste nieuwe regel komen.
                tail = first.getparent()[first.getparent().index(first) + 2 * (len(lines) - 1)]
                _set_text(tail, (tail.text or "") + suffix)
        else:
            _set_text(first, prefix + lines[0])
            for middle in nodes[start_node + 1:end_node]:
                _set_text(middle, "")
            _set_text(last, suffix)
            if len(lines) > 1:
                _insert_replacement(first, replacement)


def _all_paragraphs(document):
    """Inclusief geneste tabellen en alle unieke kop- en voetteksten."""
    seen_parts = set()

    def from_container(container):
        for paragraph in container.paragraphs:
            yield paragraph
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from from_container(cell)

    yield from from_container(document)
    for section in document.sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            if part._element in seen_parts:
                continue
            seen_parts.add(part._element)
            yield from from_container(part)


def generate_docx(template: Path, destination: Path, values: dict[str, str]) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"Word-sjabloon ontbreekt: {template}")
    if template.resolve() == destination.resolve():
        raise ValueError("Het originele Word-sjabloon mag niet worden overschreven.")
    doc = Document(str(template))
    unknown = set()
    for paragraph in _all_paragraphs(doc):
        _replace_in_paragraph(paragraph, values, unknown)
    if unknown:
        names = ", ".join("{{" + name + "}}" for name in sorted(unknown))
        raise TemplateError("Onbekende placeholders in het sjabloon: " + names)
    doc.save(str(destination))


def export_pdf(docx_path: Path, pdf_path: Path) -> tuple[bool, str]:
    """COM draait uitsluitend op Windows met Microsoft Word en pywin32."""
    try:
        import pythoncom
        from win32com.client import DispatchEx
    except ImportError:
        return False, "Microsoft Word of de PDF-module (pywin32) is niet beschikbaar."

    word = None
    opened = None
    pythoncom.CoInitialize()
    try:
        word = DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
        opened.ExportAsFixedFormat(str(pdf_path.resolve()), 17)  # wdExportFormatPDF
        return True, ""
    except Exception as exc:
        return False, f"PDF-conversie met Microsoft Word is mislukt: {exc}"
    finally:
        try:
            if opened is not None:
                opened.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
